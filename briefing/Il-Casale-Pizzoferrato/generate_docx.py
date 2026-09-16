#!/usr/bin/env python3
"""Genera la versione Word del briefing interno FC Studio — Il Casale."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

INK = RGBColor(0x1A, 0x2E, 0x24)
ACCENT = RGBColor(0x6B, 0x2E, 0x2E)
MID = RGBColor(0x8B, 0x5E, 0x3C)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB0, 0x3A, 0x2E)
CREAM = "F4F1EA"
GREEN = "1A2E24"
ROW_ALT = "F7F5F0"


def set_run_font(run, name="Calibri", size=10, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="DDD4C4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, bold=False, size=8.5, color=INK, fill=None, center=False):
    if fill:
        shade(cell, fill)
    set_cell_borders(cell)
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = 1  # center-ish


def add_heading(doc, n, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "6B2E2E")
    pBdr.append(left)
    pPr.append(pBdr)
    run = p.add_run(f"{n}. {title}")
    set_run_font(run, name="Georgia", size=14, bold=True, color=INK)


def add_sub(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_run_font(run, name="Georgia", size=11, bold=True, color=RGBColor(0x2F, 0x4A, 0x3A))


def add_p(doc, text, *, size=10, italic=False, space_after=4, bold=False, color=INK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, bold=bold, color=color)
    return p


def add_bullets(doc, items, size=10):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size)


def table(doc, headers, rows, col_widths=None, header_fill=GREEN):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, bold=True, size=8, color=WHITE, fill=header_fill)
    for r_i, row in enumerate(rows):
        fill = ROW_ALT if r_i % 2 else "FFFFFF"
        for c_i, val in enumerate(row):
            first = c_i == 0
            cell_text(
                t.rows[r_i + 1].cells[c_i],
                val,
                bold=first,
                size=8,
                color=INK,
                fill=fill,
            )
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.6)

    header = section.header
    hp = header.paragraphs[0]
    run = hp.add_run("FC Studio  ·  Scheda di briefing operativo  ·  uso interno  ·  settembre 2026")
    set_run_font(run, size=8, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Il Casale, Pizzoferrato  ·  riunione soci")
    set_run_font(run, size=8, color=MUTED)

    # Title
    k = doc.add_paragraph()
    k.paragraph_format.space_after = Pt(2)
    run = k.add_run("SCHEDA DI BRIEFING OPERATIVO  ·  USO INTERNO")
    set_run_font(run, size=9, bold=True, color=ACCENT)

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(0)
    run = t.add_run("Il Casale, Pizzoferrato")
    set_run_font(run, name="Georgia", size=22, bold=True, color=INK)

    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(6)
    run = s.add_run("Cosa sta succedendo online e cosa può fare FC Studio")
    set_run_font(run, name="Georgia", size=12, color=RGBColor(0x3D, 0x5A, 0x4C))

    meta = doc.add_paragraph()
    run = meta.add_run("Riunione soci  ·  16 settembre 2026  ·  Lettura: ~5 minuti")
    set_run_font(run, size=9, color=MUTED)

    add_p(
        doc,
        "Legenda:  Fatto osservato  ·  Interpretazione  ·  Proposta FC Studio  ·  NON VERIFICATO = non accessibile dal nostro ambiente (scheda Google nativa, feed Instagram, contenuti Facebook).",
        size=8.5,
        italic=True,
        color=MUTED,
    )

    # 1
    add_heading(doc, 1, "Snapshot del locale")
    snapshot = [
        ("Nome / luogo.", "Il Casale — Casale Pollice 4, 66040 Pizzoferrato (CH), Parco Nazionale della Majella. Online compare anche come «Ristorante Il Casale» e «Agriturismo Il Casale»."),
        ("Tipologia.", "Ristorante di montagna a formula fissa (antipasti caldi/freddi → primi a scelta → brace → dolci e liquori della casa). Non è un agriturismo in senso stretto: nessuna evidenza di produzione propria."),
        ("Posizionamento.", "Di fatto è il «pranzo della domenica» della Valle del Sangro: abbondanza, vista, prezzo percepito basso. La proposta è chiara nei clienti, assente nei canali del locale."),
        ("Target.", "Famiglie e gruppi della costa (Lanciano–Pescara–Vasto, 60–90 min) e villeggianti di Valle del Sole. Potenziale non coperto: turisti Majella/Roccaraso ed eventi privati."),
        ("Distintivi.", "Vetrata sulla vallata · tavola degli antipasti · ravioli noci e tartufo, pappardelle al cinghiale, gnocchi con orapi · titolare Carmine in sala · rapporto qualità/prezzo."),
        ("Digitale.", "Reputazione da leader locale (Google 4,6 / 661; Tripadvisor 4,4 / 239, 1° su 13). Canali propri: deboli o assenti. Il problema non è farsi conoscere: è rendersi trovabili, comprensibili e coerenti."),
    ]
    for k, v in snapshot:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(k + " ")
        set_run_font(r, size=10, bold=True)
        r = p.add_run(v)
        set_run_font(r, size=10)

    # 2
    add_heading(doc, 2, "Stato attuale")
    table(
        doc,
        ["Area", "Stato attuale", "Criticità", "Opportunità"],
        [
            [
                "Google / Maps",
                "Scheda attiva. 661 recensioni, media 4,6. Crescita recente (~70 recensioni in pochi mesi, via aggregatori). Telefono 0872 946242.",
                "Orari solo sab 12–23 / dom 12–16 (un cliente ha trovato aperto il martedì). Nessun sito né menu collegato. Attributo «consegna a domicilio» probabilmente errato.",
                "Orari stagionali, categorie (abruzzese / griglia / eventi), Q&A, post, foto proprietario. Impatto immediato su chiamate e coperti.",
            ],
            [
                "Instagram",
                "NON VERIFICATO. Nessun handle pubblico rintracciabile. I competitor strutturati (Nido del Falco) lo hanno linkato ovunque.",
                "Zero visibilità sul pubblico 25–45 e sui turisti che cercano «Abruzzo / Majella».",
                "Profilo ufficiale + Reel su vista e piatti-firma. Secondario rispetto a Google, utile per crescita.",
            ],
            [
                "Facebook",
                "Pagina esistente (ID legacy + 62 voti 5/5 via Restaurant Guru). Bio, frequenza, engagement: NON VERIFICATI.",
                "Probabile pagina luogo non gestita. Il target 45–65 della provincia è ancora qui.",
                "Piattaforma social primaria per il pubblico attuale. Eventi, Capodanno, orari del weekend.",
            ],
            [
                "Tripadvisor",
                "239 recensioni, 4,4/5, 1° su 13 a Pizzoferrato, 76 foto. Categorie: italiana, barbecue, mediterranea. Fascia «€».",
                "Nessun sito collegato. Fascia «economico» in tensione con conti recenti da 30–40 €. Stato rivendicazione e risposte: NON VERIFICATO.",
                "Rivendicare, aggiornare fascia, collegare sito. Rilevante per turisti, secondario rispetto a Google.",
            ],
            [
                "Sito web",
                "Assente. ilcasale4.placeweb.site (sito automatico Google, chiuso 2024) è parcheggiato. ilcasalepizzoferrato.it è libero (NXDOMAIN) ma ancora linkato da aggregatori.",
                "Nessuna fonte autorevole. Casa del Pastore vince la SERP «ristorante Pizzoferrato» con un Joomla datato. Tre aggregatori pubblicano dati sbagliati.",
                "Registrare il dominio già «atteso». Sito utile, ma non è il primo intervento: prima vanno sistemati Google e le informazioni.",
            ],
            [
                "Menu digitale",
                "Nessun menu ufficiale. Solo ricostruzioni di terzi (PDF incoerente: pizza, pesce). Prezzo percepito 20–40 €, nessuno ufficiale.",
                "La formula fissa, se non spiegata, produce attrito (rifiuto antipasti, incertezza sul conto).",
                "Pagina «Come funziona il pranzo» + prezzo + QR. Alto impatto, lavoro contenuto.",
            ],
            [
                "Fotografia",
                "Solo foto utente (76 TA, 162 aggregate). Autentiche, disordinate, senza brand.",
                "Nessuna immagine che presenti il locale come vuole il titolare. Google e social vivono di scatti casuali.",
                "1 giornata di shooting (sala, vetrata, 12 piatti, staff) alimenta tutti i canali per 6–12 mesi.",
            ],
            [
                "Video / content",
                "Nessun Reel, TikTok o YouTube rintracciabile associato al locale.",
                "Vista e tavola degli antipasti — i due motivi per cui si va — non sono usati come contenuto.",
                "Serie «La finestra del Casale» + time-lapse antipasti. Da fare dopo lo shooting, non prima.",
            ],
            [
                "Branding",
                "Logo: NON VERIFICATO. Nome in 4 varianti. Nessuna palette/tono ufficiale. Accoglienza fisica curata (guida rossa «BENVENUTO»).",
                "«Il Casale» è un nome omonimo di decine di locali. Online il brand è deciso dai clienti.",
                "Naming operativo «Il Casale Pizzoferrato» + kit minimo (logo, colori, avatar coerenti).",
            ],
            [
                "Local SEO",
                "Forte nel local pack per volume recensioni. Assente nell’organico (nessuna pagina propria).",
                "NAP incoerente: Casale Pollice vs Casale Greci; telefono milanese su Sluurpy; «cucina cinese» su Friends-and-food; PagineGialle non rivendicata.",
                "Bonifica directory (tempo, poco denaro). Query locali senza contenuti di ristoratori.",
            ],
            [
                "Recensioni",
                "Capitale reputazionale più grande della zona. Pregi ripetuti: antipasti, vista, Q/P, pasta in casa, staff. ~10% del campione critico sui primi elaborati.",
                "Risposte del titolare: NON VERIFICATE (nessuna visibile negli estratti). Gap di aspettativa sul prezzo. Nessun QR di sollecito.",
                "Protocollo risposte + QR sul conto. Le 900 recensioni sono già copy per social e sito.",
            ],
        ],
        col_widths=[3.2, 5.2, 5.2, 4.4],
    )

    # 3
    add_heading(doc, 3, "Punti di forza (già in mano al locale)")
    strengths = [
        ("1. Reputazione da leader", "661 recensioni Google a 4,6 + 1° su Tripadvisor.", "È la prova sociale più forte della zona; il passaparola digitale c’è già.", "Badge, citazioni, risposte che ripetono i piatti-firma."),
        ("2. La vetrata sulla vallata", "La parola «vista / vallata» è tra le più citate.", "È l’unico asset visivo non copiabile dai competitor di paese.", "Foto hero, serie «La finestra del Casale», copertine Google/social."),
        ("3. Formula abbondante a prezzo basso percepito", "Percorso completo, clienti «quasi a scoppiare», conto che sorprende in positivo.", "Promessa semplice e memorabile; spiega perché si sceglie Il Casale.", "Dichiararla online (oggi esiste solo nelle recensioni)."),
        ("4. Piatti-firma spontanei", "Ravioli noci e tartufo, pappardelle al cinghiale, gnocchi con orapi, liquore al kiwi.", "I clienti hanno già nominato i «prodotti eroe».", "Foto ufficiali, Reel, menu digitale, post stagionali."),
        ("5. Persone riconoscibili", "Carmine, Fernanda, Michele, «le ragazze» — citati per nome.", "È l’unico racconto umano disponibile; i competitor lo usano già (Antonio e Felicia al Nido del Falco).", "Pagina Chi siamo, ritratti, Stories di sala."),
        ("6. Sala + posizione", "Capienza, accessibilità, parcheggio, camino; Majella, eventi di paese, Roccaraso a ~35 min.", "Adatto a gruppi, famiglie, cerimonie, cenoni — scontrino complessivo alto.", "Pagina Eventi/Capodanno con form; contenuti geotaggati Majella."),
    ]
    for title, have, why, how in strengths:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(title)
        set_run_font(r, name="Georgia", size=11, bold=True)
        add_p(doc, have, size=10, space_after=1)
        add_p(doc, "Perché conta: " + why, size=9.5, italic=True, space_after=1, color=MUTED)
        add_p(doc, "Come usarlo: " + how, size=9.5, italic=True, space_after=6, color=MUTED)

    # 4
    add_heading(doc, 4, "Criticità a impatto commerciale")
    table(
        doc,
        ["Problema", "Conseguenza", "Soluzione"],
        [
            [
                "Orari Google solo weekend",
                "Clienti infrasettimanali persi. Già documentato da un recensore (aperto il martedì, Google diceva chiuso).",
                "Orari stagionali e festività su GBP. Costo: zero. Impatto: coperti persi recuperabili da subito.",
            ],
            [
                "Nessun menu né prezzo ufficiale",
                "Chi non conosce la formula non prenota, o arriva e rifiuta gli antipasti. Recensioni negative evitabili.",
                "Menu digitale «Come funziona» con prezzo dichiarato e QR su tavolo/ingresso.",
            ],
            [
                "Dati contraddittori in rete",
                "Due indirizzi, quattro nomi, telefono di Milano, «cucina cinese». Navigatori e Google ricevono segnali incoerenti.",
                "Bonifica NAP su 6 directory + rivendica PagineGialle/Bianche. Tempo, non budget.",
            ],
            [
                "Prenotazione solo su fisso, «obbligatoria»",
                "Telefono saturo nel weekend; nessuna traccia di richieste eventi/Capodanno.",
                "WhatsApp Business + form eventi. Riduce no-show e chiamate perse.",
            ],
            [
                "Zero canali propri (sito assente, social NON VERIFICATI / inattivi)",
                "Casa del Pastore intercetta «ristorante Pizzoferrato». Aggregatori monetizzano i suoi clienti. Nessun canale per tartufo, neve, cenoni.",
                "Prima: Google + menu. Poi: dominio già libero + Facebook (target attuale). Instagram e sito come secondo passo, non come prima vendita.",
            ],
            [
                "Nessuna foto ufficiale",
                "Chi scopre il locale su Maps non vede la vetrata come andrebbe vista. La decisione resta sulle foto dei clienti.",
                "Shooting 1 giornata. Alimenta Google, menu, sito, social per un anno.",
            ],
            [
                "Prezzo percepito vs conto reale",
                "Tripadvisor «€» e recensioni vecchie da 20–30 € vs ultime da 30–40 €. Rischio delusione sul conto, non sul cibo.",
                "Dichiarare la fascia sul menu digitale e aggiornare Tripadvisor. Governare l’aspettativa.",
            ],
        ],
        col_widths=[5.4, 6.3, 6.3],
    )
    add_p(
        doc,
        "Nota: le critiche sui primi elaborati (~10% del campione) sono un tema di cucina, non di marketing. FC Studio può farle emergere e rispondere online; non può «sistemarle» da sola.",
        size=9,
        italic=True,
        color=MUTED,
    )

    # 5
    add_heading(doc, 5, "Opportunità per FC Studio")
    add_sub(doc, "Quick wins — veloci, alto rapporto costo/impatto")
    table(
        doc,
        ["Cosa facciamo", "Problema che risolve", "Valore per il cliente"],
        [
            [
                "Ottimizzazione Google Business Profile (orari, categorie, attributi, descrizione, Q&A)",
                "Scheda incompleta / fuorviante; traffico infrasettimanale perso",
                "Più chiamate e indicazioni da chi sta già cercando un ristorante in zona",
            ],
            [
                "Bonifica NAP (6 directory + rivendica PG/PB/Tripadvisor)",
                "Indirizzo, telefono e cucina sbagliati in rete",
                "Navigatore giusto, meno clienti persi, segnali coerenti per Google",
            ],
            [
                "Menu digitale «Come funziona» + QR + prezzo",
                "Formula e conto ignoti prima di arrivare",
                "Meno attrito, meno recensioni evitabili, meno telefonate «quanto costa?»",
            ],
            [
                "WhatsApp Business + QR recensioni sul conto",
                "Telefono unico canale; recensioni lasciate al caso",
                "Prenotazioni tracciabili; rating difeso e accelerato",
            ],
            [
                "Registrazione ilcasalepizzoferrato.it",
                "Dominio libero, già linkato da aggregatori, a rischio snatch",
                "Presidio del nome; base per landing o sito",
            ],
        ],
        col_widths=[6.2, 5.9, 5.9],
    )

    add_sub(doc, "Progetti — più lavoro, costruiscono identità")
    table(
        doc,
        ["Cosa facciamo", "Problema che risolve", "Valore per il cliente"],
        [
            [
                "Mini brand kit + naming «Il Casale Pizzoferrato»",
                "Omonimia e incoerenza tra piattaforme",
                "Riconoscibilità minima su Google, social, biglietti, QR",
            ],
            [
                "Shooting foto + video (1 giornata)",
                "Nessuna immagine ufficiale",
                "Materiale per 6–12 mesi su tutti i canali; la vetrata diventa argomento di vendita",
            ],
            [
                "Landing o sito 5 pagine (Template V2.0) — Home, Formula, Chi siamo, Eventi, Contatti",
                "Nessuna fonte autorevole; SERP organica persa",
                "Info vere, form eventi/Capodanno, click da Google. Utile, non urgente quanto GBP+menu",
            ],
        ],
        col_widths=[6.2, 5.9, 5.9],
    )

    add_sub(doc, "Servizi continuativi — rapporto mensile")
    table(
        doc,
        ["Cosa facciamo", "Problema che risolve", "Valore per il cliente"],
        [
            [
                "Social light: Facebook (primario) + Instagram — 2 Reel + 2 post/sett., Stories nei weekend",
                "Nessun canale per stagioni, eventi, tartufo, neve",
                "Ritorno dei clienti abituali; scoperta da 25–45 e turisti",
            ],
            [
                "Gestione recensioni (risposte 48 h + report)",
                "Critiche sui primi senza replica; capitale reputazionale non presidiato",
                "Difesa del 4,6; keyword rinforzate nelle risposte",
            ],
            [
                "Manutenzione orari/menu + 1 post GBP/settimana",
                "Informazioni che tornano a divergere",
                "Scheda sempre vera; il titolare non deve imparare Google",
            ],
            [
                "Ads Meta solo su ponti/festività/Capodanno (opzionale)",
                "Eventi affidati al passaparola",
                "Coperti extra nelle date che già riempiono la sala. Non Google Ads: la domanda locale è già sulla scheda.",
            ],
        ],
        col_widths=[6.2, 5.9, 5.9],
    )

    # 6
    add_heading(doc, 6, "Cosa vendere")
    add_p(
        doc,
        "Non vendere «un sito». Il primo valore è sistemare le informazioni su Google e spiegare la formula. Il sito è il passo 2, se il titolare vuole anche identità e eventi.",
        size=10.5,
        bold=True,
    )

    add_sub(doc, "PACCHETTO 1 — BASE  (da proporre per primi)")
    add_p(doc, "Comprende: ottimizzazione completa GBP; bonifica 6 directory + rivendica TA/PG; menu digitale + QR + prezzo; WhatsApp Business + QR recensioni; registrazione dominio; landing one-page (orari, formula, mappa, WhatsApp).")
    add_p(doc, "Obiettivo: stoppare le perdite — orari sbagliati, dati errati, formula opaca.")
    add_p(doc, "Prezzo: 1.200 – 1.800 € una tantum  +  25 – 40 €/mese hosting e micro-aggiornamenti.", bold=True)

    add_sub(doc, "PACCHETTO 2 — PRO  (upsell naturale)")
    add_p(doc, "Comprende BASE, più: mini brand kit; shooting foto+video 1 giornata; sito 5 pagine (Template V2.0) + EN essenziale; riordino Facebook + apertura Instagram; caricamento foto/menu su Google.")
    add_p(doc, "Obiettivo: dare un volto ufficiale a una reputazione che oggi parlano solo i clienti.")
    add_p(doc, "Prezzo: 3.500 – 4.800 € setup  (o +2.200–3.000 € se BASE è già stato fatto).", bold=True)

    add_sub(doc, "PACCHETTO 3 — CONTINUATIVO  (dopo 30–60 giorni)")
    add_p(doc, "Mensile: 2 Reel + 2 post/settimana; Stories nei giorni di apertura; 1 post GBP/settimana; risposte recensioni entro 48 h; aggiornamento orari/menu; report mensile KPI.")
    add_p(doc, "Obiettivo: stagioni, eventi, ritorno. Ha senso solo se BASE (o PRO) è a posto e c’è materiale da pubblicare.")
    add_p(doc, "Prezzo: 450 – 650 €/mese su 12 mesi. Ads a parte (100–250 €/campagna, budget cliente).", bold=True)

    add_p(
        doc,
        "Prezzi IVA esclusa, allineati a piccola agenzia/freelance in Italia. Per un locale aperto principalmente nel weekend, BASE è la proposta coerente al primo incontro; PRO se il titolare vuole anche identità; il continuativo solo se c’è chi gira contenuto il sabato. Non proporre TheFork, SEO nazionale, delivery, TikTok da subito.",
        size=9,
        italic=True,
        color=MUTED,
    )

    # 7
    add_heading(doc, 7, "Priorità")
    table(
        doc,
        ["Priorità", "Intervento", "Impatto", "Difficoltà"],
        [
            ["1", "Correggere orari, attributi e categorie su Google Business Profile", "Alto", "Bassa"],
            ["2", "Menu digitale «Come funziona» con prezzo e QR", "Alto", "Bassa"],
            ["3", "Bonifica NAP / directory e rivendica delle schede", "Medio", "Bassa"],
            ["4", "WhatsApp Business + QR recensioni sul conto", "Medio", "Bassa"],
            ["5", "Shooting + (landing o sito) + Facebook/Instagram — solo dopo i punti 1–4", "Alto", "Media"],
        ],
        col_widths=[2.2, 11.6, 2.2, 2.2],
    )
    add_p(
        doc,
        "Ordine per impatto × fattibilità. Il sito non è in cima: senza Google e menu a posto, un sito bello non recupera i coperti persi per orari sbagliati.",
        size=9,
        italic=True,
        color=MUTED,
    )

    # 8
    add_heading(doc, 8, "Roadmap (se accettano)")
    add_sub(doc, "Fase 1 — 0–30 giorni · Sistemare")
    add_bullets(
        doc,
        [
            "Incontro: orari reali, prezzo formula, accessi FB/IG, anno di apertura",
            "GBP completo",
            "Bonifica directory",
            "Dominio registrato",
            "Menu digitale + QR stampati",
            "WhatsApp + protocollo recensioni",
            "Landing online",
        ],
    )
    add_sub(doc, "Fase 2 — 30–60 giorni · Costruire")
    add_bullets(
        doc,
        [
            "Brand kit minimo",
            "Shooting 1 giornata (meglio un sabato di servizio)",
            "Sito 5 pagine (se PRO) o arricchimento landing",
            "FB/IG riordinati",
            "Avvio 2 Reel + 2 post/sett. (se continuativo)",
            "Pagina Eventi / Capodanno",
        ],
    )
    add_sub(doc, "Fase 3 — 60–90 giorni · Misurare")
    add_bullets(
        doc,
        [
            "Confronto KPI vs giorno 0 (GBP Insights)",
            "Quale Reel converte: vista vs piatti vs persone",
            "Eventuale ads su ponte/Capodanno",
            "Report e decisione: restare su BASE o salire",
        ],
    )
    add_p(
        doc,
        "KPI da leggere, non da promettere a caso: visualizzazioni scheda, chiamate, indicazioni, click al sito/menu, nuove recensioni/mese (baseline stimata 15–25), % risposte, invii form eventi, click WhatsApp. I coperti li misura solo il titolare.",
        size=9,
        italic=True,
        color=MUTED,
    )

    # 9
    add_heading(doc, 9, "Cosa dobbiamo discutere noi due — 10 domande")
    questions = [
        "Cosa sappiamo fare internamente, cosa no? GBP, sito (abbiamo il Template V2.0), menu/QR: sì. Shooting e montaggio Reel: li facciamo noi o serve un fotografo esterno? Il costo cambia il prezzo del PRO.",
        "Chi pubblica le Stories il sabato e la domenica? Il locale è aperto nel weekend. Senza qualcuno in sala (titolare con cellulare, o noi in loco 1 volta al mese) il continuativo è fragile.",
        "Quante ore/mese abbiamo davvero se prende il continuativo (2 Reel + 2 post + recensioni + 1 mezza giornata in loco)? Se non tornano, non venderlo.",
        "Qual è il nostro costo pieno e il margine minimo? Prima di quotare 1.500 € o 4.000 €, quanto ci costa BASE e quanto PRO (giornate × tariffa nostra + fotografo + stampa QR + dominio).",
        "BASE come piede in porta, o PRO unico? Un titolare di ristorante di montagna che non ha mai investito in digitale accetta più facilmente 1.500 € visibili in 30 giorni che 4.000 €. L’upsell al PRO dopo lo shooting «di prova» è più credibile.",
        "Cosa promettiamo, per iscritto? Possiamo promettere: orari corretti, menu online, dati puliti, materiale fotografico, report KPI. Non possiamo promettere «più coperti» o «primi su Google Maps». Il 4,6 c’è già: non venderlo come risultato nostro.",
        "Cosa mostriamo al primo incontro? Non abbiamo (ancora) un case locale. Opzioni: mockup del Casale sul Template V2.0 + questa scheda + 3 screenshot dei dati errati (cinese, telefono Milano, dominio morto). È più forte di un portfolio generico.",
        "Cosa dobbiamo verificare col titolare prima di quotare? Orari infrasettimanali reali · accessi Facebook/Instagram · prezzo attuale della formula · se risponde già alle recensioni · se gli eventi/Capodanno gli interessano · budget massimo mentale.",
        "Chi parla col titolare e con quale tono? Non «ti serve un sito». Tre fatti: Google ti dà chiuso 5 giorni su 7; in rete risulti anche cinese e con un telefono di Milano; Casa del Pastore vince la ricerca col nome del paese. Poi i 661 voti a 4,6 che nessuno sta usando.",
        "Come ci differenziamo da un webmaster da 500 €? Non vendiamo pagine: sistemiamo il percorso da Maps al tavolo (orari, formula, prenotazione, recensioni). Il sito è uno strumento, non l’offerta. Se non siamo d’accordo su questo, rischiamo di quotare il pezzo sbagliato.",
    ]
    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"{i}.  ")
        set_run_font(r, size=11, bold=True, color=INK)
        r = p.add_run(q)
        set_run_font(r, size=10.5)

    # 10
    add_heading(doc, 10, "Executive summary")
    execs = [
        ("1. Dove si trova oggi, digitalmente?", "Leader di reputazione, ultimo di presidio. Google e Tripadvisor lo tengono in cima a Pizzoferrato; il locale non ha sito, non ha menu, non ha (per quanto verificabile) social attivi. Gli aggregatori raccontano al posto suo, spesso male."),
        ("2. Punti di forza principali?", "661 recensioni a 4,6 e 1° su Tripadvisor; vetrata; formula abbondante a prezzo percepito basso; piatti-firma già nominati dai clienti; titolare e staff riconoscibili; sala adatta a gruppi ed eventi in un territorio (Majella) che porta gente da fuori."),
        ("3. Problemi più importanti?", "Orari Google incompleti (coperti persi, già documentati). Formula e prezzo non dichiarati. Dati contraddittori in rete. Un solo canale di prenotazione (fisso). Nessuna immagine ufficiale. Nome omonimo. Questi punti bruciano soldi; un sito carente, da solo, no."),
        ("4. Dove può intervenire FC Studio?", "Prima sulle perdite: GBP, directory, menu/QR, WhatsApp. Poi sull’identità: foto, brand, sito, Facebook. Infine sulla crescita: contenuti e, solo su date clou, ads. Non TheFork, non SEO Italia, non delivery."),
        ("5. Primo passo concreto, oggi?", "Decidere internamente se proponiamo BASE (1.200–1.800 €) come prima offerta. Poi: una visita in loco di domenica — vedere sala, vetrata, orari veri, se esiste Instagram — e presentarsi al titolare con tre fatti verificati, non con un preventivo di sito."),
    ]
    for k, v in execs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(k + " ")
        set_run_font(r, size=10.5, bold=True, color=ACCENT)
        r = p.add_run(v)
        set_run_font(r, size=10.5)

    add_p(
        doc,
        "Fonti: aggregatori pubblici (Restaurant Guru agg. 25/08/2026, Sluurpy, Tripadvisor .cn, OneItalia, PagineGialle/Bianche, DNS/RDAP). Non accessibili dal nostro ambiente: scheda Google nativa (post, Q&A, foto proprietario), feed Instagram, contenuti Facebook. Qualsiasi preventivo va confermato dopo il sopralluogo.",
        size=8.5,
        italic=True,
        color=MUTED,
    )

    out = "/workspace/briefing/Il-Casale-Pizzoferrato/FC-Studio_Briefing_Il-Casale-Pizzoferrato.docx"
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build()
